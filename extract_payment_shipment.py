#!/usr/bin/env python3
"""
Extract "Payable to MEGA" amounts and "Release Shipment" broker info
from long documents (PDF, DOCX, TXT, XLSX).

Usage:
    python3 extract_payment_shipment.py <file_or_folder> [--out results.csv]
"""

import re
import sys
import csv
import json
import argparse
import subprocess
from pathlib import Path


# ---------------------------------------------------------------------------
# Dependency bootstrap
# ---------------------------------------------------------------------------

def _ensure(packages: list[str]) -> None:
    for pkg in packages:
        try:
            __import__(pkg.split("[")[0].replace("-", "_"))
        except ImportError:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", pkg])


_ensure(["pdfplumber", "python-docx", "openpyxl"])


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def extract_text_txt(path: Path) -> str:
    return path.read_text(errors="replace")


def extract_text_pdf(path: Path) -> str:
    import pdfplumber
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            # also pull table cells so amounts inside tables aren't missed
            for table in (page.extract_tables() or []):
                for row in table:
                    text += " | ".join(str(c or "") for c in row) + "\n"
            pages.append(text)
    return "\n".join(pages)


def extract_text_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            lines.append(" | ".join(str(c) if c is not None else "" for c in row))
    return "\n".join(lines)


_EXTRACTORS = {
    ".txt":  extract_text_txt,
    ".csv":  extract_text_txt,
    ".pdf":  extract_text_pdf,
    ".docx": extract_text_docx,
    ".xlsx": extract_text_xlsx,
    ".xls":  extract_text_xlsx,
}


def get_text(path: Path) -> str:
    fn = _EXTRACTORS.get(path.suffix.lower())
    if fn is None:
        print(f"[skip] unsupported file type: {path}")
        return ""
    return fn(path)


# ---------------------------------------------------------------------------
# Pattern matching
# ---------------------------------------------------------------------------

# Amount pattern: optional currency symbol / code, then digits with commas/dots
_AMOUNT = r"""
    (?:USD|CNY|RMB|HKD|\$|¥|HK\$)?\s*   # optional currency prefix
    [\d,]+(?:\.\d{1,2})?                  # digits, optional decimal
    (?:\s*(?:USD|CNY|RMB|HKD))?           # optional currency suffix
"""

# "Payable to MEGA" – captures the amount on the same line or next token
_PAY_MEGA_RE = re.compile(
    r"""
    (?:payable\s+to|pay\s+to|payment\s+to|应付)\s+
    (?:mega[\w\s.,&()\-]*?)    # MEGA + optional company suffix (dots, & etc.)
    [:\s,，]*
    (?P<amount>""" + _AMOUNT + r""")
    """,
    re.IGNORECASE | re.VERBOSE,
)

# Also catch lines like "MEGA    USD 12,345.00" inside tables
_MEGA_LINE_RE = re.compile(
    r"""
    ^[^\n]*\bmega\b[^\n]*?
    (?P<amount>""" + _AMOUNT + r""")
    [^\n]*$
    """,
    re.IGNORECASE | re.VERBOSE | re.MULTILINE,
)

# "Release shipment" – captures the broker name that follows
_SHIP_RE = re.compile(
    r"""
    (?:release\s+shipment\s+to   # "Release Shipment to:"
      |release\s+shipment        # "Release Shipment"
      |release\s+to\s+broker     # "Release to Broker"
      |release\s+to              # "Release to"
      |assigned?\s+broker        # "Assigned Broker" / "Assign Broker"
      |broker\s*[:：]            # "Broker:"
      |放行\s*(?:货物|shipment)?\s*(?:至|to)?  # Chinese "release cargo to"
    )
    [:\s,，]*
    (?P<broker>[A-Za-z\u4e00-\u9fff][\w\s\u4e00-\u9fff\-&.,()]{1,80}?)
    (?=\s*[\n|;,，。]|$)
    """,
    re.IGNORECASE | re.VERBOSE | re.MULTILINE,
)

# Looser fallback: if the above miss, grab broker after "release" on same line
_SHIP_FALLBACK_RE = re.compile(
    r"release[^\n]{0,30}?(?:to|broker|agent)[:\s]+(?P<broker>[A-Z][\w &.,\-]{2,60})",
    re.IGNORECASE,
)


def _clean_amount(raw: str) -> str:
    return re.sub(r"\s+", " ", raw).strip()


def _clean_broker(raw: str) -> str:
    return re.sub(r"\s+", " ", raw).strip().rstrip(".,;")


def parse_document(text: str) -> list[dict]:
    results = []

    for m in _PAY_MEGA_RE.finditer(text):
        amt = _clean_amount(m.group("amount"))
        context = _get_context(text, m.start(), m.end())
        broker = _find_nearby_broker(text, m.start())
        results.append({
            "type": "payable_to_mega",
            "amount": amt,
            "broker": broker,
            "context": context,
        })

    # Fallback: table rows that contain MEGA + amount but weren't caught above
    seen_starts = {m.start() for m in _PAY_MEGA_RE.finditer(text)}
    for m in _MEGA_LINE_RE.finditer(text):
        if m.start() not in seen_starts and "payable" not in m.group(0).lower():
            amt = _clean_amount(m.group("amount"))
            broker = _find_nearby_broker(text, m.start())
            results.append({
                "type": "mega_amount_line",
                "amount": amt,
                "broker": broker,
                "context": m.group(0).strip(),
            })

    # Standalone release-shipment hits (no MEGA amount nearby)
    for m in _SHIP_RE.finditer(text):
        broker = _clean_broker(m.group("broker"))
        if not broker:
            continue
        # skip if already captured as part of a mega result
        if not _near_mega(text, m.start()):
            results.append({
                "type": "release_shipment",
                "amount": "",
                "broker": broker,
                "context": _get_context(text, m.start(), m.end()),
            })

    return results


def _get_context(text: str, start: int, end: int, window: int = 120) -> str:
    left = max(0, start - window)
    right = min(len(text), end + window)
    snippet = text[left:right].replace("\n", " ")
    return re.sub(r"\s+", " ", snippet).strip()


def _near_mega(text: str, pos: int, window: int = 300) -> bool:
    snippet = text[max(0, pos - window): pos + window]
    return bool(re.search(r"\bmega\b", snippet, re.IGNORECASE))


def _find_nearby_broker(text: str, pos: int, window: int = 400) -> str:
    snippet = text[pos: pos + window]
    for pattern in (_SHIP_RE, _SHIP_FALLBACK_RE):
        m = pattern.search(snippet)
        if m:
            return _clean_broker(m.group("broker"))
    return ""


# ---------------------------------------------------------------------------
# File / folder scanning
# ---------------------------------------------------------------------------

def scan(target: Path) -> list[dict]:
    files = sorted(target.rglob("*")) if target.is_dir() else [target]
    all_rows = []
    for f in files:
        if f.suffix.lower() not in _EXTRACTORS:
            continue
        print(f"[read] {f}")
        text = get_text(f)
        if not text:
            continue
        rows = parse_document(text)
        for r in rows:
            r["file"] = str(f)
        all_rows.extend(rows)
    return all_rows


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------

def write_csv(rows: list[dict], out: Path) -> None:
    fields = ["file", "type", "amount", "broker", "context"]
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    print(f"[done] {len(rows)} row(s) → {out}")


def write_json(rows: list[dict], out: Path) -> None:
    out.write_text(json.dumps(rows, ensure_ascii=False, indent=2))
    print(f"[done] {len(rows)} row(s) → {out}")


def print_table(rows: list[dict]) -> None:
    if not rows:
        print("No matches found.")
        return
    print(f"\n{'File':<30} {'Type':<22} {'Amount':<20} {'Broker'}")
    print("-" * 100)
    for r in rows:
        fname = Path(r.get("file", "")).name
        print(f"{fname:<30} {r['type']:<22} {r['amount']:<20} {r['broker']}")
    print()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("target", help="File or folder to scan")
    ap.add_argument("--out", default="results.csv",
                    help="Output file (.csv or .json).  Default: results.csv")
    args = ap.parse_args()

    target = Path(args.target)
    if not target.exists():
        sys.exit(f"Error: {target} does not exist")

    rows = scan(target)
    print_table(rows)

    out = Path(args.out)
    if out.suffix.lower() == ".json":
        write_json(rows, out)
    else:
        write_csv(rows, out)


if __name__ == "__main__":
    main()
